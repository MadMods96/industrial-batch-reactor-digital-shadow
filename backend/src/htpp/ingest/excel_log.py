"""Excel batch log parser driven by column_map.yaml.

Plant workbooks often have Date + Reactor label but no panel batch number.
Those rows are matched to `batches` by machine and Asia/Kolkata calendar day.
Customer-specific reactor labels come from HTPP_MACHINE_ALIASES or column_map.local.yaml.
"""

from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

import duckdb
import yaml
from openpyxl import load_workbook

from htpp.config import settings
from htpp.schemas import BatchLogEntry

MAP_PATH = Path(__file__).with_name("column_map.yaml")
LOCAL_MAP_PATH = Path(__file__).with_name("column_map.local.yaml")


class ExcelMapError(RuntimeError):
    pass


def _norm(value: object) -> str:
    return " ".join(str(value or "").split()).lower()


def load_column_map() -> dict:
    data = yaml.safe_load(MAP_PATH.read_text(encoding="utf-8")) or {}
    if LOCAL_MAP_PATH.exists():
        local = yaml.safe_load(LOCAL_MAP_PATH.read_text(encoding="utf-8")) or {}
        if "machine_name_to_id" in local:
            merged = dict(data.get("machine_name_to_id") or {})
            merged.update(local["machine_name_to_id"])
            data["machine_name_to_id"] = merged
        for key, value in local.items():
            if key != "machine_name_to_id":
                data[key] = value
    aliases = settings.machine_alias_map()
    if aliases:
        merged = dict(data.get("machine_name_to_id") or {})
        merged.update(aliases)
        data["machine_name_to_id"] = merged
    return data


def parse_workbook(path: Path, column_map: dict | None = None) -> list[BatchLogEntry]:
    column_map = column_map or load_column_map()
    book = load_workbook(path, data_only=True)
    sheet_name = column_map.get("sheet") or book.sheetnames[0]
    if sheet_name not in book.sheetnames:
        # Fall back to the first sheet when the map names a missing title.
        sheet_name = book.sheetnames[0]
    sheet = book[sheet_name]
    preferred = int(column_map.get("header_row") or 1)
    header_row, headers, chosen = _pick_header(sheet, column_map["columns"], preferred)
    found = [h for h in headers if h]
    required = [
        "machine_id", "log_date", "feed_mass_kg", "moisture_pct",
        "oil_mass_kg", "carbon_mass_kg", "steel_mass_kg",
    ]
    missing = [name for name in required if name not in chosen]
    if missing:
        raise ExcelMapError(f"Required columns {missing} not found. Headers present: {found}")

    names = {_norm(name): mid for name, mid in column_map.get("machine_name_to_id", {}).items()}
    entries: list[BatchLogEntry] = []
    for offset, row in enumerate(sheet.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
        if all(cell is None or str(cell).strip() == "" for cell in row):
            continue

        def val(field: str):
            idx = chosen.get(field)
            return None if idx is None else row[idx]

        machine_id = _resolve_machine(val("machine_id"), names)
        batch_raw = val("batch_no")
        batch_no = int(float(batch_raw)) if batch_raw not in (None, "") else 0
        log_date = _parse_date(val("log_date"))
        feed = float(val("feed_mass_kg"))
        moisture = _moisture_pct(val("moisture_pct"), feed)
        shred = _optional_float(val("shred_kg"))
        mix = _optional_float(val("mix_kg"))
        feedstock = val("feedstock_type")
        if feedstock in (None, ""):
            feedstock = "tyre" if (shred or 0) >= (mix or 0) else "mix"
        mix_pct = _optional_float(val("feedstock_mix_pct"))
        if mix_pct is None and mix is not None and feed > 0:
            mix_pct = 100.0 * mix / feed
        entries.append(
            BatchLogEntry(
                machine_id=machine_id,
                batch_no=batch_no,
                log_date=log_date,
                feed_mass_kg=feed,
                moisture_pct=moisture,
                feedstock_type=str(feedstock),
                feedstock_mix_pct=mix_pct,
                oil_mass_kg=float(val("oil_mass_kg")),
                carbon_mass_kg=float(val("carbon_mass_kg")),
                steel_mass_kg=float(val("steel_mass_kg")),
                gas_mass_kg=_optional_float(val("gas_mass_kg")),
                operator=None if val("operator") is None else str(val("operator")),
                notes=None if val("notes") is None else str(val("notes")),
                source_file=path.name,
                source_row=offset,
            )
        )
    return entries


def resolve_batch_numbers(
    con: duckdb.DuckDBPyConnection,
    entries: list[BatchLogEntry],
) -> tuple[list[BatchLogEntry], list[dict]]:
    """Attach panel batch_no by machine + IST calendar day. Keep explicit batch_no as-is."""
    matched: list[BatchLogEntry] = []
    unmatched: list[dict] = []
    for entry in entries:
        if entry.batch_no and entry.batch_no > 0:
            matched.append(entry)
            continue
        if entry.log_date is None:
            unmatched.append({
                "machine_id": entry.machine_id,
                "batch_no": None,
                "log_date": None,
                "reason": "excel row has no date",
                "source_row": entry.source_row,
            })
            continue
        day = entry.log_date.date() if isinstance(entry.log_date, datetime) else entry.log_date
        row = con.execute(
            """
            SELECT batch_no, total_duration_min
            FROM batches
            WHERE machine_id = ?
              AND CAST(started_at AT TIME ZONE 'Asia/Kolkata' AS DATE) = ?
            ORDER BY total_duration_min DESC NULLS LAST, batch_no
            LIMIT 1
            """,
            [entry.machine_id, day],
        ).fetchone()
        if row is None:
            # Overnight starts: excel date often equals the end calendar day.
            row = con.execute(
                """
                SELECT batch_no, total_duration_min
                FROM batches
                WHERE machine_id = ?
                  AND CAST(ended_at AT TIME ZONE 'Asia/Kolkata' AS DATE) = ?
                ORDER BY total_duration_min DESC NULLS LAST, batch_no
                LIMIT 1
                """,
                [entry.machine_id, day],
            ).fetchone()
        if row is None:
            # Loading day can sit inside a multi-day batch window.
            row = con.execute(
                """
                SELECT batch_no, total_duration_min
                FROM batches
                WHERE machine_id = ?
                  AND CAST(started_at AT TIME ZONE 'Asia/Kolkata' AS DATE) <= ?
                  AND CAST(ended_at AT TIME ZONE 'Asia/Kolkata' AS DATE) >= ?
                ORDER BY total_duration_min DESC NULLS LAST, batch_no
                LIMIT 1
                """,
                [entry.machine_id, day, day],
            ).fetchone()
        if row is None:
            unmatched.append({
                "machine_id": entry.machine_id,
                "batch_no": None,
                "log_date": day.isoformat(),
                "reason": "no panel batch with this machine and date",
                "source_row": entry.source_row,
            })
            continue
        matched.append(entry.model_copy(update={"batch_no": int(row[0])}))
    return matched, unmatched


def _pick_header(sheet, columns: dict, preferred: int) -> tuple[int, list[str], dict[str, int]]:
    best: tuple[int, list[str], dict[str, int]] | None = None
    for row_no in range(1, min(6, sheet.max_row + 1)):
        headers = [_norm(cell.value) for cell in sheet[row_no]]
        chosen: dict[str, int] = {}
        for field, candidates in columns.items():
            for candidate in candidates:
                key = _norm(candidate)
                if key in headers:
                    chosen[field] = headers.index(key)
                    break
        score = len(chosen)
        if score >= 6 and (best is None or score > len(best[2]) or (score == len(best[2]) and row_no == preferred)):
            best = (row_no, headers, chosen)
        if row_no == preferred and score >= 6:
            return row_no, headers, chosen
    if best is None:
        headers = [_norm(cell.value) for cell in sheet[preferred]]
        return preferred, headers, {}
    return best


def _parse_date(value: object) -> datetime:
    if hasattr(value, "isoformat") and not isinstance(value, str):
        if isinstance(value, datetime):
            return value
        if isinstance(value, date):
            return datetime.combine(value, datetime.min.time())
    return datetime.fromisoformat(str(value)[:10])


def _moisture_pct(raw: object, feed_mass_kg: float) -> float:
    """Accept percent, fraction (0–1), or net moisture mass in kg."""
    value = float(raw)
    if value > 100 and feed_mass_kg > 0:
        return 100.0 * value / feed_mass_kg
    if 0 <= value <= 1.0:
        return value * 100.0
    return value


def _optional_float(value: object) -> float | None:
    if value is None or str(value).strip() == "":
        return None
    return float(value)


def _resolve_machine(raw: object, names: dict[str, int]) -> int:
    if isinstance(raw, (int, float)) and not isinstance(raw, bool):
        return int(raw)
    text = str(raw).strip()
    if text.isdigit():
        return int(text)
    key = _norm(text)
    if key not in names:
        raise ExcelMapError(
            f"Unresolvable machine name {text!r}. Known names: {sorted(names)}"
        )
    return names[key]
