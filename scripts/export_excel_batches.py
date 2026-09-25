"""Export anonymized plant Excel → plant_batches.json (Fix Brief 002/003).

Duration = sum of phase hour columns × 60 (ignore misleading "total in min").
Keys = R1-YYYY-MM-DD (no assumed panel machine-id mapping in the key).
"""

from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

SRC = Path(r"c:\Users\cto\Desktop\GA\fly\sagar_bhai_wagholi-plant.xlsx")
OUT = Path(__file__).resolve().parents[1] / "frontend" / "lib" / "demo" / "plant_batches.json"

# Display slot only — not confirmed panel join (Brief 002/003 B2).
REACTOR_SLOT = {"r1": 0, "r2": 1, "r3": 2}
REACTOR_LABEL = {"r1": "R1", "r2": "R2", "r3": "R3"}
# Provisional ids for UI grouping only; keys do not embed these.
PROVISIONAL_ID = {"r1": 1093, "r2": 1094, "r3": 1146}

# Hour columns that sum to batch duration (Brief 003 B1).
DURATION_HOURS = ["gas", "down", "cooling", "carbon", "nitrogen", "loading", "start"]


def _norm(value: object) -> str:
    return " ".join(str(value or "").split()).lower()


def main() -> None:
    wb = load_workbook(SRC, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    header_idx = None
    headers: list[str] = []
    for i, row in enumerate(rows[:5]):
        vals = [_norm(c) for c in row]
        if "date" in vals and "reactor" in vals:
            header_idx = i
            headers = vals
            break
    if header_idx is None:
        raise SystemExit("header not found")

    def col(*names: str) -> int | None:
        for name in names:
            if name in headers:
                return headers.index(name)
        return None

    i_date = col("date")
    i_reac = col("reactor")
    i_load = col("loading")  # mass loading — careful: also a duration col name
    # Mass "Loading" is at index of exact "loading" that is feed — sheet has Loading (mass) and Loading (hours)
    # From earlier dump: Loading is feed at col 10, later "Loading " is duration.
    # Prefer feed from column named exactly after Net Moisture block: loading near shred.
    i_feed = None
    for idx, name in enumerate(headers):
        if name == "loading" and i_feed is None:
            # First "loading" in this sheet is feed mass (kg).
            i_feed = idx
            break
    i_moist_net = col("net moisture")
    i_oil = col("oil(kg)")
    i_carb = col("carbon(kg)")
    i_steel = col("steel(kg)")
    i_oilp = col("oil %")
    i_total_typo = col("total in min")
    i_total_h = col("total")

    dur_idx: dict[str, int] = {}
    # Duration hour columns appear after oil % block; take last occurrence of each name.
    for idx, name in enumerate(headers):
        bare = name.strip()
        if bare in DURATION_HOURS:
            dur_idx[bare] = idx
        # sheet headers may be "gas " etc already normalized
        for key in DURATION_HOURS:
            if bare == key or bare.startswith(key):
                dur_idx[key] = idx

    batches: list[dict] = []
    for row in rows[header_idx + 1 :]:
        if not row or row[i_date] is None or row[i_reac] is None:
            continue
        reac_raw = _norm(row[i_reac])
        label = None
        for key, lab in REACTOR_LABEL.items():
            if key == reac_raw or key in reac_raw:
                label = lab
                break
        if label is None:
            continue
        slot_key = label.lower()
        mid = PROVISIONAL_ID[slot_key]

        d = row[i_date]
        day = d.date().isoformat() if isinstance(d, datetime) else str(d)[:10]

        def num(idx: int | None) -> float | None:
            if idx is None:
                return None
            value = row[idx]
            try:
                return float(value) if value is not None and str(value).strip() != "" else None
            except (TypeError, ValueError):
                return None

        hours = []
        missing_dur = False
        for key in DURATION_HOURS:
            value = num(dur_idx.get(key))
            if value is None:
                missing_dur = True
                hours.append(0.0)
            else:
                hours.append(value)
        sum_hours = sum(hours)
        duration_min = round(sum_hours * 60, 1) if not missing_dur else None

        typo_total = num(i_total_typo)
        flags: list[str] = ["carbon_imputed", "closure_not_measurable"]
        if missing_dur:
            flags.append("excel_durations_missing")
        if (
            duration_min is not None
            and typo_total is not None
            and abs(typo_total - sum_hours) > 0.05
        ):
            flags.append("excel_total_mismatch")

        feed = num(i_feed)
        oil_kg = num(i_oil)
        oil_pct = num(i_oilp)
        if oil_pct is not None and oil_pct <= 1.5:
            oil_pct *= 100
        if oil_pct is None and feed and oil_kg is not None:
            oil_pct = 100.0 * oil_kg / feed

        carbon_kg = num(i_carb)
        steel_kg = num(i_steel)
        # Carbon is plant estimate (Brief 002) — still store, flag already set.
        carbon_pct = (100.0 * carbon_kg / feed) if feed and carbon_kg is not None else None
        steel_pct = (100.0 * steel_kg / feed) if feed and steel_kg is not None else None
        water_kg = num(i_moist_net)  # output water recovered, not feed moisture

        # Chronological held-out: last ~20% per reactor
        batches.append(
            {
                "batch_key": f"{label}-{day}",
                "reactor": label,
                "machine_id": mid,
                "batch_no": None,
                "log_date": day,
                "total_duration_min": duration_min,
                "duration_hours_sum": None if missing_dur else round(sum_hours, 2),
                "peak_tr_c": None,
                "had_fault": None,
                "usable_for_training": None,
                "oil_yield_pct": round(oil_pct, 2) if oil_pct is not None else None,
                "carbon_yield_pct": round(carbon_pct, 2) if carbon_pct is not None else None,
                "steel_yield_pct": round(steel_pct, 2) if steel_pct is not None else None,
                "feed_mass_kg": feed,
                "water_recovered_kg": water_kg,
                "oil_mass_kg": oil_kg,
                "carbon_mass_kg": carbon_kg,
                "steel_mass_kg": steel_kg,
                "quality_flags": flags,
                "source": "plant_excel_log",
            }
        )

    # Assign train/holdout by reactor chronological order (doc 05 style, 80/20).
    by_reac: dict[str, list[dict]] = {}
    for batch in batches:
        by_reac.setdefault(batch["reactor"], []).append(batch)
    for reac, items in by_reac.items():
        items.sort(key=lambda b: b["log_date"])
        cut = max(1, int(len(items) * 0.8))
        for i, batch in enumerate(items):
            in_train = i < cut
            batch["in_training_split"] = in_train
            batch["usable_for_training"] = in_train and "excel_durations_missing" not in batch["quality_flags"]

    OUT.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "source": "anonymized plant batch log (mass / yield / time)",
        "note": (
            "Duration = sum(Gas,Down,Cooling,Carbon,Nitrogen,Loading,start) hours × 60. "
            "Carbon is a plant estimate. Net moisture column is water recovered (output), not feed moisture. "
            "batch_key = Reactor-YYYY-MM-DD; machine_id is provisional UI grouping only."
        ),
        "n_batches": len(batches),
        "batches": batches,
    }
    OUT.write_text(json.dumps(payload, indent=2), encoding="utf-8")
    print(f"wrote {len(batches)} batches -> {OUT}")
    r1 = [b for b in batches if b["reactor"] == "R1" and b["log_date"] == "2026-09-20"]
    if r1:
        print("09-20 R1 duration_min", r1[0]["total_duration_min"], "flags", r1[0]["quality_flags"])


if __name__ == "__main__":
    main()
